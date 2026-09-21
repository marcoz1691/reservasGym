"""Generate Word doc for William Apple/Google account guide."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE = Path(__file__).resolve().parent
md_path = BASE / "guia-william-cuentas-apple-google.md"
out_path = BASE / "guia-william-cuentas-apple-google.docx"

md = md_path.read_text(encoding="utf-8")
doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.85)
ns = doc.styles["Normal"]
ns.font.name = "Calibri"
ns.font.size = Pt(11)


def set_bg(cell, hexfill: str) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hexfill}"/>')
    )


def add_runs(p, text: str, bold=False, size=None, color=None) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            r = p.add_run(part)
            if bold:
                r.bold = True
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color


table_buf: list[str] = []
in_code = False
code_buf: list[str] = []


def flush_table() -> None:
    global table_buf
    if not table_buf:
        return
    rows: list[list[str]] = []
    for line in table_buf:
        if re.match(r"^\|?\s*[-:]+", line):
            continue
        rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    table_buf = []
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = t.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            val = row[ci] if ci < len(row) else ""
            add_runs(p, val, bold=(ri == 0), size=Pt(10))
            if ri == 0:
                set_bg(cell, "1E3A5F")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()


for line in md.splitlines():
    s = line.rstrip()
    if s.strip().startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code_buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            code_buf = []
            in_code = False
        else:
            flush_table()
            in_code = True
        continue
    if in_code:
        code_buf.append(s)
        continue
    if s.strip().startswith("|"):
        table_buf.append(s)
        continue
    flush_table()
    if not s.strip():
        continue
    if s.startswith("# "):
        p = doc.add_paragraph()
        add_runs(p, s[2:], bold=True, size=Pt(16), color=RGBColor(15, 23, 42))
    elif s.startswith("## "):
        p = doc.add_paragraph()
        add_runs(p, s[3:], bold=True, size=Pt(13), color=RGBColor(30, 58, 138))
    elif s.startswith("### "):
        p = doc.add_paragraph()
        add_runs(p, s[4:], bold=True, size=Pt(11.5), color=RGBColor(51, 65, 85))
    elif s.startswith("#### "):
        p = doc.add_paragraph()
        add_runs(p, s[5:], bold=True, size=Pt(11))
    elif s.startswith("---"):
        p = doc.add_paragraph()
        r = p.add_run("─" * 48)
        r.font.color.rgb = RGBColor(203, 213, 225)
        r.font.size = Pt(8)
    elif s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, s[2:])
    elif re.match(r"^\d+\. ", s.strip()):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\d+\. ", "", s.strip()))
    elif s.startswith("> "):
        p = doc.add_paragraph()
        add_runs(p, s[2:])
        for run in p.runs:
            run.italic = True
    else:
        p = doc.add_paragraph()
        add_runs(p, s)

flush_table()
doc.save(str(out_path))
print(f"Generated: {out_path}")
